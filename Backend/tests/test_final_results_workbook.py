from io import BytesIO
import unittest

import pandas as pd
from openpyxl import load_workbook

from Backend.services.admission_mapping.admission_workbook import build_workbook
from Backend.services.final_results_workbook import (
    FINAL_RESULT_SHEETS,
    build_final_results_workbook,
    final_results_available,
    final_results_filename,
)


class FinalResultsWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.state = {}
        for index, (stage, filename, sheet_name) in enumerate(FINAL_RESULT_SHEETS):
            rows = pd.DataFrame({f'field_{index}': [f'{sheet_name} value']})
            self.state.setdefault(stage, {})[filename] = build_workbook(rows, 'Result')

    def test_keeps_each_result_schema_in_a_separate_sheet(self):
        data = build_final_results_workbook(self.state)
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False)
        self.assertEqual(workbook.sheetnames, [item[2] for item in FINAL_RESULT_SHEETS])
        for index, (_, _, sheet_name) in enumerate(FINAL_RESULT_SHEETS):
            sheet = workbook[sheet_name]
            self.assertEqual(sheet.cell(1, 1).value, f'field_{index}')
            self.assertEqual(sheet.cell(2, 1).value, f'{sheet_name} value')

    def test_includes_available_stages_and_requires_at_least_one(self):
        self.assertTrue(final_results_available(self.state))
        admission_only = {'admission_exports': self.state['admission_exports']}
        data = build_final_results_workbook(admission_only)
        workbook = load_workbook(BytesIO(data), read_only=True)
        self.assertEqual(workbook.sheetnames, [
            'Admission Matched', 'Admission Review', 'Admission Not Matched',
        ])
        self.state.clear()
        self.assertFalse(final_results_available(self.state))
        with self.assertRaisesRegex(ValueError, 'Run at least one mapping'):
            build_final_results_workbook(self.state)

    def test_filename_uses_sanitized_school_identity(self):
        state = {
            'saved_admission_dump': {'school_index': '914', 'school_name': 'North/West: School'},
        }
        self.assertEqual(final_results_filename(state), 'automation-914-North_West_ School.xlsx')


if __name__ == '__main__':
    unittest.main()
