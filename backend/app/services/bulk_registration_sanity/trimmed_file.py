"""Trim all input records once, before publishing the stored binary snapshot."""
import csv
from io import BytesIO, StringIO
from pathlib import Path

from openpyxl import load_workbook


def trim_input_file(name, data):
    """Keep headers and workbook structure; trim data cells on every worksheet."""
    if Path(name).suffix.lower() == '.csv':
        rows = csv.reader(StringIO(data.decode('utf-8-sig')))
        output = StringIO(newline='')
        writer = csv.writer(output)
        writer.writerow(next(rows))
        writer.writerows([value.strip() for value in row] for row in rows)
        return output.getvalue().encode('utf-8-sig')
    workbook = load_workbook(BytesIO(data))
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, str):
                        original_type = cell.data_type
                        cell.value = cell.value.strip()
                        cell.data_type = original_type
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()
