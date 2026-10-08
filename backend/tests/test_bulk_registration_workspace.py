"""Workspace concurrency and committed-output regressions, using isolated files."""
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from threading import Event
import asyncio
import json
import unittest
from unittest.mock import patch

from fastapi import HTTPException, Request, Response, UploadFile
from openpyxl import Workbook, load_workbook
from sqlalchemy.exc import SQLAlchemyError

from app.routes.bulk_registration import conversion_routes, file_routes, workspace_routes
from app.routes.bulk_registration.models import StoredFileInput
from app.routes.bulk_registration.workspace_access import BULK_COOKIE
from app.services import bulk_registration_storage as storage
from app.services import operation_log
from app.main import create_app
from tests.test_file_workflow_api import asgi_request


class BulkWorkspaceTests(unittest.TestCase):
    def setUp(self):
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        root = patch.object(storage, 'ROOT', Path(temporary.name))
        root.start()
        self.addCleanup(root.stop)
        log_temporary = TemporaryDirectory()
        self.addCleanup(log_temporary.cleanup)
        log_path = patch.object(operation_log, 'LOG_PATH', Path(log_temporary.name) / 'operations.jsonl')
        log_path.start()
        self.addCleanup(log_path.stop)
        self.workspace_id = storage.create_workspace()

    def request(self):
        return Request({'type': 'http', 'headers': [
            (b'cookie', f'{BULK_COOKIE}={self.workspace_id}'.encode()),
        ]})

    def test_upload_persists_trimmed_input_and_sanity_is_read_only(self):
        uploaded = file_routes.upload_workspace_file(
            self.request(), Response(), 0,
            UploadFile(filename='students.csv', file=BytesIO(b'FIRST NAME,EMAIL\n Ada , a@example.com \n Bob , A@example.com \n')),
        )
        before = storage.load_workspace(self.workspace_id)
        result = file_routes.sanity_check(self.request(), Response(), uploaded['revision'])
        self.assertEqual(uploaded['file']['rows'][0], ['Ada', 'a@example.com'])
        self.assertEqual(result['checks'][-1]['failed_count'], 2)
        saved = storage.load_workspace(self.workspace_id)
        self.assertEqual(saved, before)
        self.assertIsNone(saved['output'])
        self.assertEqual(saved['outputs'], {})
        _, data = storage.read_snapshot(self.workspace_id, saved['input'])
        self.assertIn(b'Ada,a@example.com', data)
        self.assertNotIn(b' Ada ', data)
        self.assertEqual(file_routes.get_input_page(self.request(), Response())['file']['rows'][0],
                         ['Ada', 'a@example.com'])
        with self.assertRaises(HTTPException) as error:
            file_routes.sanity_check(self.request(), Response(), 0)
        self.assertEqual(error.exception.status_code, 409)

    def test_sanity_checks_house_against_dump_and_preserves_workspace(self):
        uploaded = file_routes.upload_workspace_file(
            self.request(), Response(), 0,
            UploadFile(filename='students.csv', file=BytesIO(
                b'FIRST NAME,house\nAda, Blue \nBob,Red\nCara,\n')),
        )
        before = storage.load_workspace(self.workspace_id)
        with patch.object(file_routes, 'fetch_houses', return_value=['Blue']) as fetch:
            result = file_routes.sanity_check(self.request(), Response(), uploaded['revision'])
        fetch.assert_called_once_with()
        self.assertEqual(result['checks'][-1], {
            'id': 'house', 'label': 'House exists in dump', 'failed_count': 1, 'passed': False,
        })
        self.assertIn('House not found in dump; insert the house', result['rows'][1]['problems'])
        self.assertEqual(storage.load_workspace(self.workspace_id), before)
        with patch.object(file_routes, 'fetch_houses', side_effect=SQLAlchemyError('unavailable')):
            with self.assertRaises(HTTPException) as error:
                file_routes.sanity_check(self.request(), Response(), uploaded['revision'])
        self.assertEqual(error.exception.status_code, 503)

    def test_sanity_without_house_does_not_query_dump(self):
        uploaded = file_routes.upload_workspace_file(
            self.request(), Response(), 0,
            UploadFile(filename='students.csv', file=BytesIO(b'FIRST NAME\nAda\n')),
        )
        with patch.object(file_routes, 'fetch_houses') as fetch:
            file_routes.sanity_check(self.request(), Response(), uploaded['revision'])
        fetch.assert_not_called()

    def test_preview_replaces_houses_with_ids_and_persists_them_for_exports(self):
        from app.services.bulk_registration import export_frame

        uploaded = file_routes.upload_workspace_file(
            self.request(), Response(), 0,
            UploadFile(filename='students.csv', file=BytesIO(
                b'FIRST NAME,House\nBob, BLUE \nAda,green\nCara,\nDora,Unknown\n')),
        )
        with patch.object(conversion_routes, 'get_school', return_value={
            'school_index': '914', 'school_name': 'Test School',
        }), patch.object(conversion_routes, 'fetch_sections', return_value=[]), patch.object(
            conversion_routes, 'fetch_house_ids', return_value=[(12, 'Blue'), (37, 'Green')],
        ) as fetch, patch.object(
            conversion_routes, 'fetch_available_usernames',
            return_value=['ada001', 'bob001', 'cara001', 'dora001'],
        ):
            result = conversion_routes.convert_workspace_file(
                self.request(), Response(), uploaded['revision'], '914', 'preview',
                sheet=None, page=1)
        fetch.assert_called_once_with()
        column = result['output']['columns'].index('house')
        expected = ['37', '12', '', 'Unknown']
        self.assertEqual([row[column] for row in result['output']['rows']], expected)
        state = storage.load_workspace(self.workspace_id)
        _, output = conversion_routes.read_authoritative_output(self.workspace_id, state)
        self.assertEqual(output['house'].tolist(), expected)
        self.assertEqual(output.attrs['unresolved_house_positions'], [3])
        page = conversion_routes.get_output_page(self.request(), Response())
        self.assertEqual([row[column] for row in page['output']['rows']], expected)
        workbook = load_workbook(BytesIO(export_frame(output, 'xlsx')), read_only=True)
        self.assertEqual(
            [row[column] or '' for row in workbook.active.iter_rows(min_row=2, values_only=True)],
            expected,
        )
        workbook.close()
        _, source = storage.read_snapshot(self.workspace_id, state['input'])
        self.assertIn(b'Bob,BLUE', source)
        self.assertIn(b'Ada,green', source)

    def test_house_validation_blocks_exports_and_rechecks_database(self):
        for house in ('Unknown', '12', 'Blue', ''):
            with self.subTest(house=house):
                self.workspace_id = storage.create_workspace()
                uploaded = file_routes.upload_workspace_file(
                    self.request(), Response(), 0,
                    UploadFile(filename='students.csv', file=BytesIO(
                        ('FIRST NAME,FULL NAME,Section,Class Number,GENDER,house\n'
                         f'Bob,Bob Smith,A,Class I,Male,{house}\n').encode())),
                )
                with patch.object(conversion_routes, 'get_school', return_value={
                    'school_index': '914', 'school_name': 'Test School',
                }), patch.object(conversion_routes, 'fetch_sections', return_value=[(1, 'A')]), patch.object(
                    conversion_routes, 'fetch_house_ids', return_value=[(12, 'Blue')],
                ) as houses, patch.object(
                    conversion_routes, 'fetch_available_usernames', return_value=['bob001'],
                ), patch.object(conversion_routes, 'fetch_existing_usernames', return_value=set()), patch.object(
                    conversion_routes, 'fetch_existing_emails', return_value=set(),
                ):
                    preview = conversion_routes.convert_workspace_file(
                        self.request(), Response(), uploaded['revision'], '914', 'preview',
                        sheet=None, page=1)
                    verified = conversion_routes.verify_bulk_registration_output(
                        self.request(), Response(), preview['revision'])
                    expected = house in ('Blue', '')
                    self.assertEqual(verified['verification']['passed'], expected)
                    revision = verified['workspace']['revision']
                    if expected:
                        exported = conversion_routes.convert_workspace_file(
                            self.request(), Response(), revision, '914', 'csv',
                            sheet=None, page=1)
                        self.assertIn(b'bob001', exported.body)
                        if not house:
                            houses.assert_not_called()
                            continue
                        houses.return_value = []
                    else:
                        # Even a previously verified/legacy snapshot must be rechecked.
                        state = storage.load_workspace(self.workspace_id)
                        state['output_verified'] = True
                        storage.save_workspace(self.workspace_id, state)
                    with self.assertRaises(HTTPException) as error:
                        conversion_routes.convert_workspace_file(
                            self.request(), Response(), revision, '914', 'csv',
                            sheet=None, page=1)
                    self.assertEqual(error.exception.status_code, 409)
                    self.assertFalse(storage.load_workspace(self.workspace_id)['output_verified'])

    def test_trimmed_workbook_preserves_other_sheets_and_drives_generated_output(self):
        workbook = Workbook()
        workbook.active.title = 'Students'
        workbook.active.append(['FIRST NAME', 'FULL NAME', 'CONTACT'])
        workbook.active.append([' Ada ', ' Ada Lovelace ', ' 001234 '])
        workbook.create_sheet('Other').append([' Untouched '])
        workbook['Other'].append([' Trim this too '])
        source = BytesIO()
        workbook.save(source)
        workbook.close()
        result = file_routes.upload_workspace_file(self.request(), Response(), 0,
            UploadFile(filename='students.xlsx', file=BytesIO(source.getvalue())), None)
        self.assertEqual(result['file']['rows'][0], ['Ada', 'Ada Lovelace', '001234'])
        state = storage.load_workspace(self.workspace_id)
        _, data = storage.read_snapshot(self.workspace_id, state['input'])
        saved = load_workbook(BytesIO(data))
        self.assertEqual(saved['Students']['A2'].value, 'Ada')
        self.assertEqual(saved['Other']['A1'].value, ' Untouched ')
        self.assertEqual(saved['Other']['A2'].value, 'Trim this too')
        saved.close()
        with patch.object(conversion_routes, 'get_school', return_value={
            'school_index': '914', 'school_name': 'Test School',
        }), patch.object(conversion_routes, 'fetch_sections', return_value=[]), patch.object(
            conversion_routes, 'fetch_available_usernames', return_value=['ada001'],
        ):
            preview = conversion_routes.convert_workspace_file(
                self.request(), Response(), 1, '914', 'preview',
                sheet=None, page=1)
            output = preview['output']
            self.assertEqual(output['rows'][0][output['columns'].index('FIRST NAME')], 'Ada')
            with self.assertRaises(HTTPException) as unverified:
                conversion_routes.convert_workspace_file(
                    self.request(), Response(), 2, '914', 'csv',
                    sheet=None, page=1)
            self.assertEqual(unverified.exception.status_code, 409)
            self.assertIn('verify successfully', unverified.exception.detail)
            state = storage.load_workspace(self.workspace_id)
            state['output_verified'] = True
            storage.save_workspace(self.workspace_id, state)
            with patch.object(conversion_routes, 'fetch_existing_usernames', return_value=set()), patch.object(
                conversion_routes, 'fetch_existing_emails', return_value=set(),
            ), patch.object(conversion_routes, 'verify_final_output', return_value={'passed': True}):
                download = conversion_routes.convert_workspace_file(
                    self.request(), Response(), 2, '914', 'csv',
                    sheet=None, page=1)
            self.assertIn(b'Ada Lovelace', download.body)
            self.assertNotIn(b' Ada ', download.body)

    def test_stale_upload_after_reset_cannot_replace_the_cookie(self):
        app = create_app()
        old_cookie = f'{BULK_COOKIE}={self.workspace_id}'.encode()
        status, headers, body = asyncio.run(asgi_request(
            app, 'DELETE', '/api/v1/bulk-reg/workspace', headers=[
                (b'cookie', old_cookie), (b'x-workspace-revision', b'0'),
            ],
        ))
        self.assertEqual(status, 200)
        reset = json.loads(body)
        new_cookie = headers[b'set-cookie'].split(b';', 1)[0]
        folders = set(storage.ROOT.iterdir())
        multipart = (
            b'--review-boundary\r\n'
            b'Content-Disposition: form-data; name="file"; filename="students.csv"\r\n'
            b'Content-Type: text/csv\r\n\r\n'
            b'FIRST NAME\nAda\n\r\n--review-boundary--\r\n'
        )

        def upload(cookie):
            return asyncio.run(asgi_request(app, 'POST', '/api/v1/bulk-reg/files', multipart, [
                (b'cookie', cookie), (b'x-workspace-revision', b'0'),
                (b'content-type', b'multipart/form-data; boundary=review-boundary'),
            ]))

        for cookie in (old_cookie, b'', f'{BULK_COOKIE}=invalid'.encode()):
            with self.subTest(cookie=cookie):
                status, headers, _ = upload(cookie)
                self.assertEqual(status, 409)
                self.assertNotIn(b'set-cookie', headers)
                self.assertEqual(set(storage.ROOT.iterdir()), folders)

        status, _, body = asyncio.run(asgi_request(
            app, 'GET', '/api/v1/bulk-reg/workspace', headers=[(b'cookie', new_cookie)],
        ))
        self.assertEqual(status, 200)
        self.assertEqual(json.loads(body), reset)
        status, headers, body = upload(new_cookie)
        self.assertEqual(status, 200)
        self.assertNotIn(b'set-cookie', headers)
        self.assertEqual(json.loads(body)['workspace_id'], reset['workspace_id'])

    def test_workspace_initialization_still_recovers_a_missing_workspace(self):
        storage.delete_workspace(self.workspace_id)
        response = Response()
        result = workspace_routes.initialize_workspace(self.request(), response)
        self.assertIn('set-cookie', response.headers)
        self.assertEqual(result['revision'], 0)
        self.assertIsNone(result['file'])

    def test_delayed_refresh_after_reset_cannot_replace_the_cookie(self):
        app = create_app()
        old_cookie = f'{BULK_COOKIE}={self.workspace_id}'.encode()
        status, headers, body = asyncio.run(asgi_request(
            app, 'DELETE', '/api/v1/bulk-reg/workspace', headers=[
                (b'cookie', old_cookie), (b'x-workspace-revision', b'0'),
            ],
        ))
        self.assertEqual(status, 200)
        reset = json.loads(body)
        new_cookie = headers[b'set-cookie'].split(b';', 1)[0]
        folders = set(storage.ROOT.iterdir())
        for cookie in (old_cookie, b'', f'{BULK_COOKIE}=invalid'.encode()):
            with self.subTest(cookie=cookie):
                status, headers, _ = asyncio.run(asgi_request(
                    app, 'GET', '/api/v1/bulk-reg/workspace', headers=[(b'cookie', cookie)],
                ))
                self.assertEqual(status, 409)
                self.assertNotIn(b'set-cookie', headers)
                self.assertEqual(set(storage.ROOT.iterdir()), folders)
        # Initialization is idempotent when another tab has already recovered.
        status, headers, body = asyncio.run(asgi_request(
            app, 'POST', '/api/v1/bulk-reg/workspace', headers=[(b'cookie', new_cookie)],
        ))
        self.assertEqual(status, 200)
        self.assertNotIn(b'set-cookie', headers)
        self.assertEqual(json.loads(body), reset)

    def test_simultaneous_mutations_reject_the_stale_revision(self):
        first_loaded, second_started, second_loaded = Event(), Event(), Event()
        original = workspace_routes.workspace_for

        def controlled_read(request, response):
            result = original(request, response)
            if not first_loaded.is_set():
                first_loaded.set()
                self.assertTrue(second_started.wait(5))
                # Without the operation lock, the second request reads revision 0
                # here. With it, that read waits until the first request commits.
                second_loaded.wait(0.5)
            else:
                second_loaded.set()
            return result

        def clear(second=False):
            if second:
                second_started.set()
            try:
                result = workspace_routes.clear_file(self.request(), Response(), 0)
                return 200, result['revision']
            except HTTPException as error:
                return error.status_code, None

        with patch.object(workspace_routes, 'workspace_for', controlled_read), ThreadPoolExecutor(2) as pool:
            first = pool.submit(clear)
            self.assertTrue(first_loaded.wait(5))
            second = pool.submit(clear, True)
            self.assertEqual(first.result(timeout=5), (200, 1))
            self.assertEqual(second.result(timeout=5), (409, None))
        self.assertEqual(storage.load_workspace(self.workspace_id)['revision'], 1)

    def test_sheet_change_and_failed_conversion_preserve_output_and_passwords(self):
        workbook = Workbook()
        workbook.active.title = 'Students'
        workbook.active.append(['FIRST NAME'])
        workbook.active.append(['Ada'])
        workbook.create_sheet('Empty')
        source = BytesIO()
        workbook.save(source)
        workbook.close()
        file_routes.upload_workspace_file(
            self.request(), Response(), 0,
            UploadFile(filename='students.xlsx', file=BytesIO(source.getvalue())), None,
        )

        def convert(revision, format='preview'):
            return conversion_routes.convert_workspace_file(
                self.request(), Response(), revision, '914', format,
                sheet=None, page=1,
            )

        with patch.object(conversion_routes, 'get_school', return_value={
            'school_index': '914', 'school_name': 'Test School',
        }), patch.object(conversion_routes, 'fetch_sections', return_value=[]), patch.object(
            conversion_routes, 'fetch_available_usernames', return_value=['ada001'],
        ):
            preview = convert(1)
            state = storage.load_workspace(self.workspace_id)
            state['output_verified'] = True
            storage.save_workspace(self.workspace_id, state)
            with patch.object(conversion_routes, 'fetch_existing_usernames', return_value=set()), patch.object(
                conversion_routes, 'fetch_existing_emails', return_value=set(),
            ), patch.object(conversion_routes, 'verify_final_output', return_value={'passed': True}):
                before = convert(2, 'csv').body
            changed = file_routes.load_stored_file(
                StoredFileInput(sheet='Empty'), self.request(), Response(), 2,
            )
            self.assertEqual(changed['output'], preview['output'])
            self.assertEqual(changed['file']['sheet'], 'Empty')
            with self.assertRaises(HTTPException) as failure:
                convert(3)
            self.assertEqual(failure.exception.status_code, 400)
            with patch.object(conversion_routes, 'fetch_existing_usernames', return_value=set()), patch.object(
                conversion_routes, 'fetch_existing_emails', return_value=set(),
            ), patch.object(conversion_routes, 'verify_final_output', return_value={'passed': True}):
                self.assertEqual(convert(3, 'csv').body, before)
            state = storage.load_workspace(self.workspace_id)
            self.assertEqual(state['revision'], 3)
            self.assertEqual(state['output'], preview['output'])
            storage.read_snapshot(self.workspace_id, state['outputs']['authoritative'])
