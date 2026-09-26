import unittest
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch, MagicMock

import pandas as pd
from fastapi import HTTPException

from app.routes.bulk_registration import FilePathInput, load_path, preview_file
from app.routes.bulk_registration import convert_file, get_school
from app.services.bulk_registration import OUTPUT_HEADERS, FIXED_VALUES, convert_frame, export_frame, fetch_school
from fastapi import UploadFile


class BulkRegistrationTests(unittest.TestCase):
    def test_working_sheet_preview_and_conversion(self):
        data = BytesIO()
        with pd.ExcelWriter(data, engine='openpyxl') as writer:
            pd.DataFrame().to_excel(writer, sheet_name='Empty', index=False)
            pd.DataFrame({'FIRST NAME': ['Ada'], 'admission_number': ['001']}).to_excel(writer, sheet_name='Students', index=False)
        preview = preview_file('input.xlsx', data.getvalue())
        self.assertEqual(preview['sheets'], ['Empty', 'Students'])
        self.assertEqual(preview['sheet'], 'Empty')
        self.assertEqual(preview['row_count'], 0)
        selected = preview_file('input.xlsx', data.getvalue(), 'Students')
        self.assertEqual(selected['rows'], [['Ada', '001']])
        with self.assertRaises(HTTPException):
            preview_file('input.xlsx', data.getvalue(), 'Missing')
        with patch('app.routes.bulk_registration.fetch_school', return_value={'school_index': '42', 'school_name': 'Test'}):
            uploaded = UploadFile(filename='input.xlsx', file=BytesIO(data.getvalue()))
            result = convert_file('42', 'preview', uploaded, None, 'Students')
            self.assertEqual(result['rows'][0][0], 'Ada')
            with TemporaryDirectory() as directory:
                path = Path(directory) / 'input.xlsx'
                path.write_bytes(data.getvalue())
                with patch('app.routes.bulk_registration.settings.ALLOW_LOCAL_FILE_PATHS', True):
                    self.assertEqual(load_path(FilePathInput(path=str(path), sheet='Students'))['sheet'], 'Students')
                    result = convert_file('42', 'csv', None, str(path), 'Students')
                    self.assertIn(b'Ada', result.body)
                    with self.assertRaises(HTTPException):
                        convert_file('42', 'preview', None, str(path), 'Empty')

    def test_xlsx_preserves_formula_like_text_and_large_exports(self):
        from openpyxl import load_workbook
        output = convert_frame(pd.DataFrame({'FIRST NAME': ['=1+1'] * 2000, 'admission_number': ['001'] * 2000}),
                               {'school_index': '42', 'school_name': 'Test School'})
        workbook = load_workbook(BytesIO(export_frame(output, 'xlsx')), read_only=True)
        rows = list(workbook.active.iter_rows())
        self.assertEqual(len(rows), 2001)
        self.assertEqual(rows[-1][0].value, '=1+1')
        self.assertEqual(rows[-1][0].data_type, 's')
        self.assertEqual(rows[-1][20].value, '001')
        workbook.close()

    def test_invalid_source_does_not_query_database(self):
        with patch('app.routes.bulk_registration.fetch_school') as fetch:
            with self.assertRaises(HTTPException) as error:
                convert_file('42', 'preview', None, None)
            self.assertEqual(error.exception.status_code, 400)
            fetch.assert_not_called()

    def test_conversion_rules_and_column_order(self):
        frame = pd.DataFrame({'FIRST NAME': ['Ada', 'Sam'], 'admission_number': ['001', '002'],
                              'year': ['1999', ''], 'SCHOOL': ['Wrong', 'Wrong'],
                              'School Number': ['999', '999'], 'user_package': ['12', '11']})
        output = convert_frame(frame, {'school_index': '42', 'school_name': 'Test School'})
        self.assertEqual(list(output.columns), OUTPUT_HEADERS)
        self.assertEqual(len(output.columns), 24)
        self.assertEqual(output['admission_number'].tolist(), ['001', '002'])
        self.assertEqual(output['FULL NAME'].tolist(), ['', ''])
        self.assertEqual(output['SCHOOL'].tolist(), ['Test School'] * 2)
        self.assertEqual(output['School Number'].tolist(), ['42'] * 2)
        for column, value in FIXED_VALUES.items():
            self.assertEqual(output[column].tolist(), [value] * 2)
        for format in ('csv', 'xlsx'):
            data = BytesIO(export_frame(output, format))
            restored = (pd.read_csv(data, dtype=str, keep_default_na=False) if format == 'csv'
                        else pd.read_excel(data, dtype=str, keep_default_na=False))
            pd.testing.assert_frame_equal(restored, output)

    def test_unknown_header_rejected(self):
        with self.assertRaisesRegex(ValueError, 'Unrecognized input headers'):
            convert_frame(pd.DataFrame({'extra': ['value']}), {'school_index': '42', 'school_name': 'Test'})

    def test_conversion_api(self):
        data = b'FIRST NAME,admission_number\nAda,001\n'
        def uploaded():
            return UploadFile(filename='students.csv', file=BytesIO(data))
        with patch('app.routes.bulk_registration.fetch_school', return_value={'school_index': '42', 'school_name': 'My School'}):
            response = convert_file('42', 'preview', uploaded(), None)
            self.assertEqual(response['columns'], OUTPUT_HEADERS)
            self.assertEqual(response['rows'][0][13:15], ['My School', '42'])
            response = convert_file('42', 'xlsx', uploaded(), None)
            self.assertEqual(response.status_code, 200)
            self.assertIn('bulk-registration-42.xlsx', response.headers['content-disposition'])
            with TemporaryDirectory() as directory:
                path = Path(directory) / 'students.csv'
                path.write_bytes(data)
                with patch('app.routes.bulk_registration.settings.ALLOW_LOCAL_FILE_PATHS', True):
                    response = convert_file('42', 'preview', None, str(path))
                    self.assertEqual(response['rows'][0][20], '001')
                    self.assertEqual(response['rows'][0][13], 'My School')

    def test_school_verification(self):
        from sqlalchemy.exc import SQLAlchemyError
        connection = MagicMock()
        with patch('app.config.database.engine.connect') as connect:
            connect.return_value.__enter__.return_value = connection
            connection.execute.return_value.first.return_value = (' Test School ',)
            self.assertEqual(fetch_school(' 0042 '), {'school_index': '0042', 'school_name': 'Test School'})
            self.assertEqual(connection.execute.call_args.args[1], {'school_index': '0042'})
            for row in [None, (' ',), (None,)]:
                connection.execute.return_value.first.return_value = row
                with self.assertRaises(HTTPException) as error:
                    get_school('42')
                self.assertEqual(error.exception.status_code, 404 if row is None else 400)
            connect.side_effect = SQLAlchemyError('Unavailable')
            with self.assertRaises(HTTPException) as error:
                get_school('42')
            self.assertEqual(error.exception.status_code, 503)
        for index in ['', 'abc']:
            with self.assertRaises(HTTPException) as error:
                get_school(index)
            self.assertEqual(error.exception.status_code, 400)

    def test_csv_preserves_identifiers_and_limits_preview(self):
        result = preview_file('students.csv', b'id,name\n' + b'001,Student\n' * 25)
        self.assertEqual(result['row_count'], 25)
        self.assertEqual(len(result['rows']), 20)
        self.assertEqual(result['rows'][0], ['001', 'Student'])

    def test_xlsx(self):
        source = BytesIO()
        pd.DataFrame({'name': ['Student']}).to_excel(source, index=False)
        self.assertEqual(preview_file('students.xlsx', source.getvalue())['rows'], [['Student']])

    def test_invalid_files(self):
        for name, data in [('data.txt', b'text'), ('data.csv', b''),
                           ('data.csv', b'name\n'), ('data.xlsx', b'invalid')]:
            with self.subTest(name=name, data=data), self.assertRaises(HTTPException) as error:
                preview_file(name, data)
            self.assertEqual(error.exception.status_code, 400)

    def test_size_limit(self):
        with patch('app.routes.bulk_registration.settings.MAX_UPLOAD_BYTES', 2):
            with self.assertRaises(HTTPException) as error:
                preview_file('data.csv', b'name\nStudent')
            self.assertEqual(error.exception.status_code, 413)

    def test_path_loading_and_disabled_policy(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / 'students.csv'
            path.write_text('id,name\n001,Student', encoding='utf-8')
            with patch('app.routes.bulk_registration.settings.ALLOW_LOCAL_FILE_PATHS', True):
                self.assertEqual(load_path(FilePathInput(path=str(path)))['row_count'], 1)
            with patch('app.routes.bulk_registration.settings.ALLOW_LOCAL_FILE_PATHS', False):
                with self.assertRaises(HTTPException) as error:
                    load_path(FilePathInput(path=str(path)))
                self.assertEqual(error.exception.status_code, 403)
