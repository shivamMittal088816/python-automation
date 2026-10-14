"""Rollback file upload and school verification API tests."""

import asyncio
from contextlib import contextmanager
from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import MagicMock, patch

from openpyxl import Workbook, load_workbook
import pandas as pd

from app.main import create_app
from app.queries import query_source
from app.repositories.rollback_update_repository import (
    CANDIDATES_QUERY,
    FALLBACK_UPDATE_QUERY,
    LOCKED_CANDIDATES_QUERY,
    _analyze_with_connection,
    _plan,
    apply_rollback_updates,
)
from app.services import rollback as rollback_storage
from tests.test_file_workflow_api import asgi_request


class RollbackAPITests(unittest.TestCase):
    def setUp(self):
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        root = patch.object(rollback_storage.config, 'ROOT', Path(temporary.name))
        root.start()
        self.addCleanup(root.stop)
        self.app = create_app()
        self.dump_result = {
            'columns': ['user_id', 'admission_number', 'user_edu_school'],
            'rows': [['1', '0007', '914'], ['2', '0008', '914']],
        }

        @contextmanager
        def stream_dump(school_index):
            self.streamed_school_indexes.append(school_index)
            yield self.dump_result['columns'], iter(self.dump_result['rows'])

        self.streamed_school_indexes = []
        dump = patch(
            'app.routes.rollback.school_routes.stream_rollback_dump',
            side_effect=stream_dump,
        )
        self.stream_rollback_dump = dump.start()
        self.addCleanup(dump.stop)

    def upload(self, filename, data, content_type='application/octet-stream'):
        boundary = b'rollback-boundary'
        body = (
            b'--' + boundary + b'\r\n'
            + f'Content-Disposition: form-data; name="file"; filename="{filename}"\r\n'.encode()
            + f'Content-Type: {content_type}\r\n\r\n'.encode()
            + data + b'\r\n--' + boundary + b'--\r\n'
        )
        return asyncio.run(asgi_request(
            self.app, 'POST', '/api/v1/bulk-reg/rollback/files', body,
            [(b'content-type', b'multipart/form-data; boundary=rollback-boundary')],
        ))

    def test_fallback_update_is_scoped_to_school_students_package_14_and_excludes_sheet_ids(self):
        query = query_source('rollback_fallback_update.sql')

        self.assertIn('user_edu_school = :school_index', query)
        self.assertIn("user_type = '0'", query)
        self.assertIn('user_package = 14', query)
        self.assertIn('user_id NOT IN :supplied_ids', query)
        self.assertIn("user_subscribed = '0'", query)
        self.assertIn("user_activated = '0'", query)
        self.assertIn(
            'user_edu_class = CAST(TRIM(user_edu_class) AS UNSIGNED) - 1',
            query,
        )

    def test_commit_candidate_query_locks_rows(self):
        unlocked = query_source('rollback_candidates.sql')
        locked = query_source('rollback_candidates_for_update.sql')

        self.assertNotIn('FOR UPDATE', unlocked.upper())
        self.assertTrue(locked.upper().endswith('FOR UPDATE'))

    def test_analysis_selects_locking_query_only_when_requested(self):
        frame = pd.DataFrame({'user_id': ['123']})
        candidate = {
            'user_id': 123,
            'user_edu_class': '5',
            'user_package': 14,
            'user_activated': '1',
            'user_subscribed': '1',
            'user_subscription_date': '2026-04-01 00:00:00',
        }
        for lock_rows, expected_query in (
            (False, CANDIDATES_QUERY),
            (True, LOCKED_CANDIDATES_QUERY),
        ):
            with self.subTest(lock_rows=lock_rows):
                connection = MagicMock()
                streaming_connection = connection.execution_options.return_value
                streaming_connection.execute.return_value.mappings.return_value = iter([candidate])

                _analyze_with_connection(
                    connection, frame, '399', lock_rows=lock_rows,
                )

                streaming_connection.execute.assert_called_once_with(
                    expected_query, {'school_index': '399'},
                )

    @patch('app.repositories.rollback_update_repository._analyze_with_connection')
    @patch('app.config.database.engine')
    def test_commit_locks_verifies_and_updates_on_one_transaction_connection(
        self, engine, analyze,
    ):
        frame = pd.DataFrame({'user_id': ['123']})
        token = 'a' * 64
        analyze.return_value = {
            'preview_token': token,
            'supplied_matched': 1,
            'not_found_count': 0,
        }
        connection = engine.begin.return_value.__enter__.return_value
        connection.execute.return_value.rowcount = 4

        result = apply_rollback_updates(frame, '399', token)

        engine.begin.assert_called_once_with()
        engine.connect.assert_not_called()
        analyze.assert_called_once_with(
            connection, frame, '399', lock_rows=True,
        )
        connection.execute.assert_called_once_with(
            FALLBACK_UPDATE_QUERY,
            {'school_index': '399', 'supplied_ids': [123]},
        )
        self.assertEqual(result['fallback_updated'], 4)

    @patch('app.repositories.rollback_update_repository._analyze_with_connection')
    @patch('app.config.database.engine')
    def test_commit_hash_mismatch_rolls_back_without_update(self, engine, analyze):
        frame = pd.DataFrame({'user_id': ['123']})
        analyze.return_value = {'preview_token': 'b' * 64}
        transaction = engine.begin.return_value
        connection = transaction.__enter__.return_value

        with self.assertRaisesRegex(ValueError, 'Analyze the changes again'):
            apply_rollback_updates(frame, '399', 'a' * 64)

        connection.execute.assert_not_called()
        transaction.__exit__.assert_called_once()
        exception_type = transaction.__exit__.call_args.args[0]
        self.assertIs(exception_type, ValueError)

    def test_rollback_plan_matches_sheet_ids_after_integer_normalization(self):
        frame = pd.DataFrame({'user_id': ['00123']})
        database_rows = [
            {
                'user_id': 123, 'user_edu_class': 5, 'user_package': 14,
                'user_activated': 1, 'user_subscribed': 1,
                'user_subscription_date': '2026-04-01 00:00:00',
            },
            {
                'user_id': 124, 'user_edu_class': 5, 'user_package': 14,
                'user_activated': 1, 'user_subscribed': 1,
                'user_subscription_date': '2026-04-01 00:00:00',
            },
        ]

        plan = _plan(frame, 399, database_rows)

        self.assertEqual(plan['supplied_count'], 1)
        self.assertEqual(plan['supplied_matched'], 1)
        self.assertEqual(plan['fallback_count'], 1)
        self.assertEqual(plan['changed_count'], 1)
        self.assertEqual(plan['not_found_count'], 0)
        self.assertEqual(plan['missing_ids'], [])
        self.assertEqual(plan['changes'], [])

    def test_rollback_plan_rejects_unsafe_fallback_class_values(self):
        invalid_values = (None, '', 'Class 5', '1.5', '-1', '0')
        for value in invalid_values:
            with self.subTest(value=value), self.assertRaisesRegex(
                ValueError, '1 fallback package-14 user',
            ):
                _plan(
                    pd.DataFrame({'user_id': ['999']}),
                    399,
                    [{
                        'user_id': 123,
                        'user_edu_class': value,
                        'user_package': 14,
                        'user_activated': '1',
                        'user_subscribed': '1',
                        'user_subscription_date': '2026-04-01 00:00:00',
                    }],
                )

    def test_rollback_plan_accepts_numeric_fallback_class_text(self):
        plan = _plan(
            pd.DataFrame({'user_id': ['999']}),
            399,
            [{
                'user_id': 123,
                'user_edu_class': '05',
                'user_package': 14,
                'user_activated': '1',
                'user_subscribed': '1',
                'user_subscription_date': '2026-04-01 00:00:00',
            }],
        )

        self.assertEqual(plan['fallback_count'], 1)

    def test_protected_sheet_user_does_not_require_a_fallback_class(self):
        plan = _plan(
            pd.DataFrame({'user_id': ['00123']}),
            399,
            [{
                'user_id': 123,
                'user_edu_class': 'not numeric',
                'user_package': 14,
                'user_activated': '1',
                'user_subscribed': '1',
                'user_subscription_date': '2026-04-01 00:00:00',
            }],
        )

        self.assertEqual(plan['supplied_matched'], 1)
        self.assertEqual(plan['fallback_count'], 0)

    def test_rollback_token_ignores_fields_unrelated_to_fallback_update(self):
        frame = pd.DataFrame({'user_id': ['999']})
        base = {
            'user_id': 123,
            'user_edu_class': '5',
            'user_package': 14,
            'user_activated': '1',
            'user_subscribed': '1',
            'user_subscription_date': '2026-04-01 00:00:00',
            'user_email': 'before@example.com',
            'user_gender': 'Male',
            'user_edu_major': 'A',
        }
        unrelated_changes = {
            **base,
            'user_email': 'after@example.com',
            'user_gender': 'Female',
            'user_edu_major': 'B',
        }

        before = _plan(frame, 399, [base])
        after = _plan(frame, 399, [unrelated_changes])

        self.assertEqual(before['preview_token'], after['preview_token'])

    def test_rollback_token_changes_when_a_fallback_field_changes(self):
        frame = pd.DataFrame({'user_id': ['999']})
        base = {
            'user_id': 123,
            'user_edu_class': '5',
            'user_package': 14,
            'user_activated': '1',
            'user_subscribed': '1',
            'user_subscription_date': '2026-04-01 00:00:00',
        }

        before = _plan(frame, 399, [base])
        after = _plan(frame, 399, [{**base, 'user_edu_class': '6'}])

        self.assertNotEqual(before['preview_token'], after['preview_token'])

    def test_upload_accepts_browser_csv_file(self):
        data = b'id,name\n' + b'1,Ada\n' * 25
        status, _, response = self.upload('rollback.csv', data, 'text/csv')

        self.assertEqual(status, 200)
        result = json.loads(response)
        self.assertEqual(result['name'], 'rollback.csv')
        self.assertEqual(result['row_count'], 25)
        self.assertEqual(result['columns'], ['id', 'name'])
        snapshots = list(rollback_storage.root_path().glob('*/*.bin'))
        self.assertEqual(len(snapshots), 1)
        self.assertEqual(snapshots[0].read_bytes(), data)

        status, _, response = asyncio.run(asgi_request(
            self.app, 'GET', '/api/v1/bulk-reg/rollback/files/preview?page=2',
        ))
        self.assertEqual(status, 200)
        self.assertEqual(json.loads(response)['page'], 2)

        status, _, response = asyncio.run(asgi_request(
            self.app, 'GET', '/api/v1/bulk-reg/rollback/workspace',
        ))
        self.assertEqual(status, 200)
        self.assertEqual(json.loads(response)['file']['name'], 'rollback.csv')

        status, _, response = asyncio.run(asgi_request(
            self.app, 'DELETE', '/api/v1/bulk-reg/rollback/files',
        ))
        self.assertEqual(status, 200)
        self.assertIsNone(json.loads(response)['file'])
        self.assertEqual(list(rollback_storage.root_path().glob('*/*.bin')), [])
        state = json.loads(next(rollback_storage.root_path().glob('*/state.json')).read_text())
        self.assertIsNone(state['input'])

    def test_sanity_check_enforces_server_owned_exact_header_policy(self):
        self.upload(
            'rollback.csv',
            b'USER_ID,FIRST NAME,unexpected\n1,Ada,value\n',
            'text/csv',
        )

        status, _, response = asyncio.run(asgi_request(
            self.app, 'POST', '/api/v1/bulk-reg/rollback/files/sanity-check',
        ))

        self.assertEqual(status, 200)
        invalid = json.loads(response)
        self.assertFalse(invalid['valid'])
        self.assertEqual(invalid['missing_headers'], ['user_id'])
        self.assertEqual(invalid['unsupported_headers'], ['USER_ID', 'unexpected'])
        self.assertEqual(invalid['row_count'], 1)
        self.assertEqual(invalid['failed_count'], 1)
        self.assertIn(
            'user_id is required for rollback activity because it is the unique identifier',
            invalid['rows'][0]['problems'],
        )

        self.upload(
            'rollback.csv',
            b' user_id ,FIRST NAME,School Number\n1,Ada,914\n',
            'text/csv',
        )
        status, _, response = asyncio.run(asgi_request(
            self.app, 'POST', '/api/v1/bulk-reg/rollback/files/sanity-check',
        ))

        self.assertEqual(status, 200)
        valid = json.loads(response)
        self.assertTrue(valid['valid'])
        self.assertEqual(valid['missing_headers'], [])
        self.assertEqual(valid['unsupported_headers'], [])
        self.assertEqual(
            valid['message'],
            'Sanity check passed. Every rollback record has a unique user_id.',
        )
        self.assertEqual(valid['failed_count'], 0)
        self.assertTrue(all(check['passed'] for check in valid['checks']))

    def test_sanity_check_reports_blank_and_duplicate_user_ids_by_row(self):
        self.upload(
            'rollback.csv',
            b'user_id,FIRST NAME\n7,Ada\n,Grace\n7,Linus\n8,Alan\n',
            'text/csv',
        )

        status, _, response = asyncio.run(asgi_request(
            self.app, 'POST', '/api/v1/bulk-reg/rollback/files/sanity-check',
        ))

        self.assertEqual(status, 200)
        result = json.loads(response)
        self.assertFalse(result['valid'])
        self.assertEqual(result['row_count'], 4)
        self.assertEqual(result['failed_count'], 3)
        duplicate_problem = (
            'Duplicate user_id in rollback file after numeric normalization'
        )
        self.assertEqual(result['rows'][0]['problems'], [duplicate_problem])
        self.assertEqual(result['rows'][1]['problems'], ['user_id missing'])
        self.assertEqual(result['rows'][2]['problems'], [duplicate_problem])
        self.assertEqual(result['rows'][3]['problems'], [])
        checks = {check['id']: check for check in result['checks']}
        self.assertEqual(checks['user_id_present']['failed_count'], 1)
        self.assertEqual(checks['unique_user_id']['failed_count'], 2)

    def test_sanity_check_requires_unsigned_integer_user_ids(self):
        self.upload(
            'rollback.csv',
            b'user_id,FIRST NAME\nAbC,Ada\n-1,Grace\n1.5,Linus\n4294967296,Alan\n',
            'text/csv',
        )

        status, _, response = asyncio.run(asgi_request(
            self.app, 'POST', '/api/v1/bulk-reg/rollback/files/sanity-check',
        ))

        self.assertEqual(status, 200)
        result = json.loads(response)
        self.assertFalse(result['valid'])
        self.assertEqual(result['failed_count'], 4)
        for row in result['rows']:
            self.assertEqual(
                row['problems'],
                ['user_id invalid: use digits from 0 to 4294967295 only'],
            )
        checks = {check['id']: check for check in result['checks']}
        self.assertEqual(checks['user_id_unsigned_int']['failed_count'], 4)

    def test_sanity_check_detects_duplicates_after_numeric_normalization(self):
        self.upload(
            'rollback.csv',
            b'user_id,FIRST NAME\n00123,Ada\n123,Grace\n0007,Linus\n',
            'text/csv',
        )

        status, _, response = asyncio.run(asgi_request(
            self.app, 'POST', '/api/v1/bulk-reg/rollback/files/sanity-check',
        ))

        self.assertEqual(status, 200)
        result = json.loads(response)
        self.assertFalse(result['valid'])
        self.assertEqual(result['failed_count'], 2)
        duplicate_problem = (
            'Duplicate user_id in rollback file after numeric normalization'
        )
        self.assertEqual(result['rows'][0]['problems'], [duplicate_problem])
        self.assertEqual(result['rows'][1]['problems'], [duplicate_problem])
        self.assertEqual(result['rows'][2]['problems'], [])

    def test_sanity_check_requires_letter_only_first_names_when_header_is_present(self):
        self.upload(
            'rollback.csv',
            b'user_id,FIRST NAME\n1,Ada\n2,\n3,Ada Mary\n4,Anne-Marie\n5,aDA\n',
            'text/csv',
        )

        status, _, response = asyncio.run(asgi_request(
            self.app, 'POST', '/api/v1/bulk-reg/rollback/files/sanity-check',
        ))

        self.assertEqual(status, 200)
        result = json.loads(response)
        self.assertFalse(result['valid'])
        self.assertEqual(result['failed_count'], 3)
        self.assertEqual(result['rows'][0]['problems'], [])
        self.assertEqual(result['rows'][1]['problems'], ['First name missing'])
        self.assertEqual(
            result['rows'][2]['problems'],
            ['First name: only A-Z and a-z (no spaces) allowed'],
        )
        self.assertEqual(
            result['rows'][3]['problems'],
            ['First name: only A-Z and a-z (no spaces) allowed'],
        )
        self.assertEqual(result['rows'][4]['problems'], [])
        checks = {check['id']: check for check in result['checks']}
        self.assertEqual(checks['first_name_required']['failed_count'], 1)
        self.assertEqual(checks['first_name_characters']['failed_count'], 2)

    def test_sanity_check_allows_first_name_header_to_be_omitted(self):
        self.upload('rollback.csv', b'user_id\n1\n2\n', 'text/csv')

        status, _, response = asyncio.run(asgi_request(
            self.app, 'POST', '/api/v1/bulk-reg/rollback/files/sanity-check',
        ))

        self.assertEqual(status, 200)
        result = json.loads(response)
        self.assertTrue(result['valid'])
        self.assertNotIn(
            'first_name_required',
            [check['id'] for check in result['checks']],
        )

    def test_sanity_check_allows_blank_last_name_and_restricts_supplied_characters(self):
        self.upload(
            'rollback.csv',
            b'user_id,LAST NAME\n1,\n2,de Souza\n3,Smith\n4,Anne-Marie\n5,O\'Neil\n6,Last2\n',
            'text/csv',
        )

        status, _, response = asyncio.run(asgi_request(
            self.app, 'POST', '/api/v1/bulk-reg/rollback/files/sanity-check',
        ))

        self.assertEqual(status, 200)
        result = json.loads(response)
        self.assertFalse(result['valid'])
        self.assertEqual(result['failed_count'], 3)
        self.assertEqual(result['rows'][0]['problems'], [])
        self.assertEqual(result['rows'][1]['problems'], [])
        self.assertEqual(result['rows'][2]['problems'], [])
        for position in (3, 4, 5):
            self.assertEqual(
                result['rows'][position]['problems'],
                ['Last name: only A-Z, a-z and spaces allowed'],
            )
        checks = {check['id']: check for check in result['checks']}
        self.assertEqual(checks['last_name_characters']['failed_count'], 3)

    def test_sanity_check_requires_letter_and_space_only_full_names(self):
        self.upload(
            'rollback.csv',
            b'user_id,FULL NAME\n1,Ada Lovelace\n2,\n3,Ada2 Lovelace\n4,Anne-Marie Smith\n5,aDA\n',
            'text/csv',
        )

        status, _, response = asyncio.run(asgi_request(
            self.app, 'POST', '/api/v1/bulk-reg/rollback/files/sanity-check',
        ))

        self.assertEqual(status, 200)
        result = json.loads(response)
        self.assertFalse(result['valid'])
        self.assertEqual(result['failed_count'], 3)
        self.assertEqual(result['rows'][0]['problems'], [])
        self.assertEqual(result['rows'][1]['problems'], ['Full name missing'])
        for position in (2, 3):
            self.assertEqual(
                result['rows'][position]['problems'],
                ['Full name: only A-Z, a-z and spaces allowed'],
            )
        self.assertEqual(result['rows'][4]['problems'], [])
        checks = {check['id']: check for check in result['checks']}
        self.assertEqual(checks['full_name_required']['failed_count'], 1)
        self.assertEqual(checks['full_name_characters']['failed_count'], 2)

    def test_sanity_check_requires_exact_case_sensitive_gender_values(self):
        self.upload(
            'rollback.csv',
            b'user_id,GENDER\n1,Male\n2,Female\n3,Others\n4,\n5,male\n6,Other\n',
            'text/csv',
        )

        status, _, response = asyncio.run(asgi_request(
            self.app, 'POST', '/api/v1/bulk-reg/rollback/files/sanity-check',
        ))

        self.assertEqual(status, 200)
        result = json.loads(response)
        self.assertFalse(result['valid'])
        self.assertEqual(result['failed_count'], 3)
        self.assertTrue(all(result['rows'][position]['problems'] == []
                            for position in (0, 1, 2)))
        self.assertEqual(result['rows'][3]['problems'], ['Gender missing'])
        for position in (4, 5):
            self.assertEqual(
                result['rows'][position]['problems'],
                ['Gender invalid: use exactly Male, Female or Others'],
            )
        checks = {check['id']: check for check in result['checks']}
        self.assertEqual(checks['gender_required']['failed_count'], 1)
        self.assertEqual(checks['gender_value']['failed_count'], 2)

    def test_sanity_check_requires_fixed_subscription_values_when_headers_are_present(self):
        self.upload(
            'rollback.csv',
            (
                b'user_id,user_subscription_date,user_package,user_activated,user_subscribed\n'
                b'1,2026-04-01 00:00:00,14,1,1\n'
                b'2,,,,\n'
                b'3,2026-04-02 00:00:00,15,0,true\n'
            ),
            'text/csv',
        )

        status, _, response = asyncio.run(asgi_request(
            self.app, 'POST', '/api/v1/bulk-reg/rollback/files/sanity-check',
        ))

        self.assertEqual(status, 200)
        result = json.loads(response)
        self.assertFalse(result['valid'])
        self.assertEqual(result['failed_count'], 2)
        self.assertEqual(result['rows'][0]['problems'], [])
        self.assertEqual(result['rows'][1]['problems'], [
            'user_subscription_date missing',
            'user_package missing',
            'user_activated missing',
            'user_subscribed missing',
        ])
        self.assertEqual(result['rows'][2]['problems'], [
            'user_subscription_date invalid: use exactly 2026-04-01 00:00:00',
            'user_package invalid: use exactly 14',
            'user_activated invalid: use exactly 1',
            'user_subscribed invalid: use exactly 1',
        ])
        checks = {check['id']: check for check in result['checks']}
        for header in (
            'user_subscription_date', 'user_package',
            'user_activated', 'user_subscribed',
        ):
            self.assertEqual(checks[f'{header}_required']['failed_count'], 1)
            self.assertEqual(checks[f'{header}_value']['failed_count'], 1)

    @patch('app.routes.rollback.execution_routes.apply_rollback_updates')
    @patch('app.routes.rollback.execution_routes.preview_rollback_updates')
    @patch('app.routes.rollback.school_routes.fetch_school')
    def test_run_applies_valid_sheet_to_verified_school(
        self, fetch_school, preview_updates, apply_updates,
    ):
        fetch_school.return_value = {'school_index': '914', 'school_name': 'North School'}
        preview_token = 'a' * 64
        preview_updates.return_value = {
            'preview_token': preview_token, 'supplied_count': 1,
            'supplied_matched': 1, 'fallback_count': 4,
            'changed_count': 5, 'missing_ids': [], 'changes': [],
            'updated_fields': [],
        }
        apply_updates.return_value = {
            'supplied_count': 1,
            'supplied_updated': 0,
            'supplied_excluded': 1,
            'not_found_count': 0,
            'fallback_updated': 4,
            'updated_fields': [
                'Section', 'GENDER', 'Class Number', 'EMAIL',
                'user_subscription_date', 'user_package',
                'user_activated', 'user_subscribed',
            ],
        }
        payload = json.dumps({'school_index': '914'}).encode()
        status, _, _ = asyncio.run(asgi_request(
            self.app, 'POST', '/api/v1/bulk-reg/rollback/school', payload,
            [(b'content-type', b'application/json')],
        ))
        self.assertEqual(status, 200)
        self.upload(
            'rollback.csv',
            (
                b'user_id,Section,GENDER,Class Number,EMAIL,user_subscription_date,'
                b'user_package,user_activated,user_subscribed\n'
                b'7,A,Male,10,student@example.com,2026-04-01 00:00:00,14,1,1\n'
            ),
            'text/csv',
        )

        status, _, preview_response = asyncio.run(asgi_request(
            self.app, 'POST', '/api/v1/bulk-reg/rollback/preview-changes',
        ))
        self.assertEqual(status, 200)
        self.assertEqual(json.loads(preview_response)['preview_token'], preview_token)

        status, _, response = asyncio.run(asgi_request(
            self.app, 'POST', '/api/v1/bulk-reg/rollback/run',
            json.dumps({'preview_token': preview_token}).encode(),
            [(b'content-type', b'application/json')],
        ))

        self.assertEqual(status, 200)
        result = json.loads(response)
        self.assertEqual(result['supplied_updated'], 0)
        self.assertEqual(result['supplied_excluded'], 1)
        self.assertEqual(result['fallback_updated'], 4)
        frame, school_index, supplied_token = apply_updates.call_args.args
        self.assertEqual(school_index, '914')
        self.assertEqual(supplied_token, preview_token)
        self.assertEqual(frame.loc[0, 'user_id'], '7')
        self.assertEqual(frame.loc[0, 'GENDER'], 'Male')

    def test_excel_upload_can_select_a_worksheet(self):
        workbook = Workbook()
        workbook.active.title = 'First'
        workbook.active.append(['name'])
        workbook.active.append(['Ada'])
        second = workbook.create_sheet('Second')
        second.append(['name'])
        second.append(['Grace'])
        source = BytesIO()
        workbook.save(source)

        status, _, _ = self.upload(
            'rollback.xlsx', source.getvalue(),
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertEqual(status, 200)
        status, _, response = asyncio.run(asgi_request(
            self.app, 'POST', '/api/v1/bulk-reg/rollback/files/sheet',
            json.dumps({'sheet': 'Second'}).encode(),
            [(b'content-type', b'application/json')],
        ))

        self.assertEqual(status, 200)
        result = json.loads(response)
        self.assertEqual(result['sheets'], ['First', 'Second'])
        self.assertEqual(result['sheet'], 'Second')
        self.assertEqual(result['rows'], [['Grace']])

    def test_legacy_json_dump_is_migrated_to_bin_on_read(self):
        workspace_id = rollback_storage.create_workspace()
        folder = rollback_storage.root_path() / workspace_id
        legacy = {
            'columns': ['user_id', 'admission_number'],
            'rows': [['1', '0007'], ['2', '0008']],
        }
        (folder / 'dump.json').write_text(json.dumps(legacy), encoding='utf-8')
        state = rollback_storage.load_state(workspace_id)
        state['school'] = {'school_index': '914', 'school_name': 'North School'}
        state['dump'] = {
            'file': 'dump.json', 'row_count': 2, 'columns': legacy['columns'],
        }
        rollback_storage.save_state(workspace_id, state)

        preview = rollback_storage.read_dump_page(workspace_id, 1, 20)

        self.assertEqual(preview['rows'], legacy['rows'])
        migrated = rollback_storage.load_state(workspace_id)
        self.assertTrue(migrated['dump']['file'].startswith('dump-'))
        self.assertTrue(migrated['dump']['file'].endswith('.bin'))
        self.assertFalse((folder / 'dump.json').exists())
        self.assertTrue((folder / migrated['dump']['file']).is_file())

    def test_dump_replacement_commits_manifest_before_deleting_previous_snapshot(self):
        workspace_id = rollback_storage.create_workspace()
        school = {'school_index': '914', 'school_name': 'North School'}
        rollback_storage.save_school_rows(
            workspace_id, school, ['user_id'], iter([['1']]),
        )
        before = rollback_storage.load_state(workspace_id)
        previous = rollback_storage.root_path() / workspace_id / before['dump']['file']

        with patch('app.services.rollback.dump_snapshots.save_state', side_effect=OSError('disk full')):
            with self.assertRaises(OSError):
                rollback_storage.save_school_rows(
                    workspace_id, school, ['user_id'], iter([['2']]),
                )

        self.assertTrue(previous.is_file())
        persisted = json.loads(
            (rollback_storage.root_path() / workspace_id / 'state.json').read_text(encoding='utf-8')
        )
        self.assertEqual(persisted['dump']['file'], before['dump']['file'])

    @patch('app.routes.bulk_registration.school_routes.fetch_school')
    def test_school_number_returns_database_name(self, fetch_school):
        fetch_school.return_value = {'school_index': '914', 'school_name': 'North School'}

        status, _, response = asyncio.run(asgi_request(
            self.app, 'GET', '/api/v1/bulk-reg/schools/914',
        ))

        self.assertEqual(status, 200)
        self.assertEqual(json.loads(response), {
            'school_index': '914', 'school_name': 'North School',
        })
        fetch_school.assert_called_once_with('914')

    @patch('app.routes.rollback.school_routes.fetch_school')
    def test_rollback_school_is_restored_and_can_be_cleared(self, fetch_school):
        fetch_school.return_value = {'school_index': '914', 'school_name': 'North School'}
        payload = json.dumps({'school_index': '914'}).encode()

        status, _, response = asyncio.run(asgi_request(
            self.app, 'POST', '/api/v1/bulk-reg/rollback/school', payload,
            [(b'content-type', b'application/json')],
        ))
        self.assertEqual(status, 200)
        self.assertEqual(json.loads(response)['school']['school_name'], 'North School')
        self.assertEqual(self.streamed_school_indexes, ['914'])

        self.upload('rollback.csv', b'id,name\n1,Ada\n', 'text/csv')
        status, _, response = asyncio.run(asgi_request(
            self.app, 'GET', '/api/v1/bulk-reg/rollback/workspace',
        ))
        self.assertEqual(status, 200)
        restored = json.loads(response)
        self.assertEqual(restored['school'], {
            'school_index': '914', 'school_name': 'North School',
        })
        self.assertEqual(restored['file']['name'], 'rollback.csv')
        self.assertEqual(restored['dump']['row_count'], 2)
        self.assertEqual(restored['dump']['rows'][0], ['1', '0007', '914'])
        folder = next(rollback_storage.root_path().iterdir())
        dump_snapshots = list(folder.glob('dump-*.bin'))
        self.assertEqual(len(dump_snapshots), 1)
        self.assertIn(b'user_id,admission_number,user_edu_school', dump_snapshots[0].read_bytes())

        status, _, response = asyncio.run(asgi_request(
            self.app, 'DELETE', '/api/v1/bulk-reg/rollback/school',
        ))
        self.assertEqual(status, 200)
        self.assertIsNone(json.loads(response)['school'])

        status, _, response = asyncio.run(asgi_request(
            self.app, 'GET', '/api/v1/bulk-reg/rollback/workspace',
        ))
        restored = json.loads(response)
        self.assertIsNone(restored['school'])
        self.assertIsNone(restored['dump'])
        self.assertEqual(restored['file']['name'], 'rollback.csv')
        self.assertEqual(list(folder.glob('dump-*.bin')), [])

    @patch('app.routes.rollback.school_routes.fetch_school')
    def test_dump_preview_is_paginated_for_verified_school(self, fetch_school):
        fetch_school.return_value = {'school_index': '399', 'school_name': 'Example School'}
        self.dump_result = {
            'columns': ['user_id', 'user_edu_school'],
            'rows': [['=2+2' if index == 0 else str(index), '399'] for index in range(25)],
        }
        payload = json.dumps({'school_index': '399'}).encode()

        status, _, response = asyncio.run(asgi_request(
            self.app, 'POST', '/api/v1/bulk-reg/rollback/school', payload,
            [(b'content-type', b'application/json')],
        ))
        result = json.loads(response)
        self.assertEqual(status, 200)
        self.assertEqual(result['dump']['row_count'], 25)
        self.assertEqual(len(result['dump']['rows']), 20)

        status, _, response = asyncio.run(asgi_request(
            self.app, 'GET', '/api/v1/bulk-reg/rollback/dump/preview?page=2',
        ))
        second = json.loads(response)
        self.assertEqual(status, 200)
        self.assertEqual(second['page'], 2)
        self.assertEqual(len(second['rows']), 5)
        self.assertEqual(second['rows'][0], ['20', '399'])
        self.assertEqual(self.streamed_school_indexes, ['399'])

        status, headers, csv_data = asyncio.run(asgi_request(
            self.app, 'GET', '/api/v1/bulk-reg/rollback/dump/download?file_format=csv',
        ))
        self.assertEqual(status, 200)
        self.assertIn(b'rollback-dump-399.csv', headers[b'content-disposition'])
        self.assertTrue(csv_data.startswith(b'\xef\xbb\xbfuser_id,user_edu_school\n'))
        self.assertIn(b"'=2+2,399\n", csv_data)
        self.assertIn(b'24,399\n', csv_data)

        status, headers, xlsx_data = asyncio.run(asgi_request(
            self.app, 'GET', '/api/v1/bulk-reg/rollback/dump/download?file_format=xlsx',
        ))
        self.assertEqual(status, 200)
        self.assertIn(b'rollback-dump-399.xlsx', headers[b'content-disposition'])
        workbook = load_workbook(BytesIO(xlsx_data), read_only=True, data_only=True)
        sheet = workbook['Rollback dump']
        self.assertEqual([cell.value for cell in next(sheet.iter_rows())],
                         ['user_id', 'user_edu_school'])
        self.assertEqual(sheet['A2'].value, '=2+2')
        self.assertEqual(sheet['A26'].value, '24')


if __name__ == '__main__':
    unittest.main()
